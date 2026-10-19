"""Source-derived invoice facts and qualification bound to an exact PDF rendition."""

import argparse
import hashlib
import inspect
import json
import re
import shutil
import unicodedata
import uuid
from collections.abc import Callable
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import Any, cast

import pymupdf
import yaml
from docx import Document

from pdf_bench.persistence import write_json_atomic

HEADER_ALIASES = {
    "Description": "description",
    "Quantity": "quantity",
    "Qty": "quantity",
    "Unit": "unit",
    "Unit Price": "unit_price",
    "Rate": "unit_price",
    "Amount": "line_total",
}
POPULATION = tuple(f"invoice_{number:03d}" for number in range(1, 101))
MONTHS = (
    "January February March April May June July August September October November December"
).split()


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def normalized_text(value: str) -> str:
    return " ".join(unicodedata.normalize("NFKC", value).split())


def decimal_string(value: str) -> str:
    value = normalized_text(value)
    negative = value.startswith("(") and value.endswith(")")
    if negative:
        value = value[1:-1]
    value = value.removeprefix("$").strip()
    if not re.fullmatch(r"[+-]?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?", value):
        raise ValueError(f"Unsupported displayed decimal: {value!r}")
    number = Decimal(value.replace(",", ""))
    if negative:
        number = number.copy_negate()
    result = format(number, "f")
    return result.rstrip("0").rstrip(".") if "." in result else result


def source_date(value: str) -> str:
    value = normalized_text(value)
    if re.fullmatch(r"\d{2}/\d{2}/\d{4}", value):
        month, day, year = map(int, value.split("/"))
    else:
        match = re.fullmatch(r"([A-Za-z]+) (\d{1,2}), (\d{4})", value)
        if match is None or match[1] not in MONTHS:
            raise ValueError(f"Unsupported source date: {value!r}")
        month, day, year = MONTHS.index(match[1]) + 1, int(match[2]), int(match[3])
    return date(year, month, day).isoformat()


def extract_reference(source: Path | bytes) -> dict:
    """Read printed source facts; do not consult metadata, old references or outputs."""
    document = Document(BytesIO(source if isinstance(source, bytes) else source.read_bytes()))
    nodes = [(f"paragraph/{i}", p.text) for i, p in enumerate(document.paragraphs)]
    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            nodes.append(
                (
                    f"table/{table_index}/row/{row_index}",
                    "\n".join(cell.text for cell in row.cells),
                )
            )

    def field(pattern: str) -> tuple[str, list[str]]:
        matches = [
            (match[1].strip(), location)
            for location, text in nodes
            for match in re.finditer(pattern, text, flags=re.IGNORECASE | re.MULTILINE)
        ]
        if not matches or len({normalized_text(value) for value, _ in matches}) != 1:
            raise ValueError(f"Missing or ambiguous source field: {pattern}")
        return matches[0][0], [location for _, location in matches]

    invoice_number, number_locations = field(r"Invoice\s*(?:#|Number)\s*:\s*([^\s|]+)")
    due, due_locations = field(
        r"(?:Due Date|Payment Due|Due)\s*:\s*"
        r"([A-Za-z]+ \d{1,2}, \d{4}|\d{2}/\d{2}/\d{4})"
    )
    total, total_locations = field(r"^TOTAL\s*:\s*(\$[\d,]+\.\d{2})\s*$")
    header = document.tables[0]
    if len(header.rows) == 1 and len(header.columns) == 2:
        vendor_lines = [
            line.strip() for line in header.cell(0, 0).text.splitlines() if line.strip()
        ]
        vendor = vendor_lines[1] if vendor_lines[0].upper() == "FROM:" else vendor_lines[0]
        vendor_location = "table/0/row/0/column/0"
    else:
        paragraphs = [(i, p.text) for i, p in enumerate(document.paragraphs) if p.text.strip()]
        vendor_location = f"paragraph/{paragraphs[1][0]}"
        vendor = paragraphs[1][1].splitlines()[0]
    items = []
    source_headers = None
    raw_records = []
    locations = []
    for table_index, table in enumerate(document.tables):
        headers = [normalized_text(cell.text) for cell in table.rows[0].cells]
        if "Description" not in headers:
            continue
        if source_headers is not None or any(h not in HEADER_ALIASES for h in headers):
            raise ValueError("Unsupported or ambiguous source line-item table")
        if len(set(headers)) != len(headers):
            raise ValueError("Duplicate source headers")
        source_headers = headers
        for row_index, row in enumerate(table.rows[1:], 1):
            surfaces = {
                HEADER_ALIASES[h]: cell.text for h, cell in zip(headers, row.cells, strict=True)
            }
            item = {
                key: decimal_string(value)
                if key in {"quantity", "unit_price", "line_total"}
                else normalized_text(value)
                for key, value in surfaces.items()
            }
            raw_records.append(surfaces)
            items.append(item)
            locations.append(f"table/{table_index}/row/{row_index}")
    if not source_headers or not items:
        raise ValueError("Source has no supported line-item facts")
    return {
        "schema_version": 1,
        "fields": {
            "invoice_number": normalized_text(invoice_number),
            "vendor": normalized_text(vendor),
            "due_date": source_date(due),
            "grand_total": decimal_string(total),
            "currency_marker": "$",
        },
        "source_headers": source_headers,
        "applicable_record_fields": [HEADER_ALIASES[h] for h in source_headers],
        "records": items,
        "surfaces": {
            "invoice_number": invoice_number,
            "vendor": vendor,
            "due_date": due,
            "grand_total": total,
            "records": raw_records,
        },
        "source_locations": {
            "invoice_number": number_locations,
            "due_date": due_locations,
            "grand_total": total_locations,
            "vendor": vendor_location,
            "records": locations,
        },
    }


def reference_protocol() -> dict:
    code = "\n".join(
        inspect.getsource(f)
        for f in cast(
            list[Callable[..., Any]],
            [normalized_text, decimal_string, source_date, extract_reference],
        )
    )
    return {
        "schema_version": 1,
        "task": "Controlled born-digital invoice to fields and complete line-item records",
        "header_aliases": HEADER_ALIASES,
        "english_month_names": MONTHS,
        "unicode_database_version": unicodedata.unidata_version,
        "numeric_date_policy": "MM/DD/YYYY, as printed by the retained synthetic generator",
        "amount_policy": "Decimal strings from displayed source values; no arithmetic correction",
        "currency_policy": "Printed dollar marker only; no inferred ISO currency",
        "source_extractor_sha256": hashlib.sha256(code.encode()).hexdigest(),
        "metadata_policy": "Cross-check only, never reference truth",
        "reference_policy": "Source truth requires independently reviewed exact PDF rendition",
        "qualification_implementation_sha256": hashlib.sha256(
            "\n".join(
                inspect.getsource(f)
                for f in cast(
                    list[Callable[..., Any]],
                    [
                        mechanical_presence,
                        validate_reference_bundle,
                        render_reference_pages,
                        qualify_references,
                    ],
                )
            ).encode()
        ).hexdigest(),
        "rendering_pymupdf_version": pymupdf.VersionBind,
        "population_scope": "Retained controlled templates; not a sample of real vendor invoices",
    }


def mechanical_presence(pdf: Path | bytes, reference: dict) -> dict:
    """Unscored text-presence diagnostic; it cannot establish visible associations."""
    data = pdf if isinstance(pdf, bytes) else pdf.read_bytes()
    with pymupdf.open(stream=data, filetype="pdf") as document:
        pages = [normalized_text(page.get_text()) for page in document]
    text = " ".join(pages)
    values = {key: value for key, value in reference["surfaces"].items() if key != "records"}
    for index, row in enumerate(reference["surfaces"]["records"]):
        values.update({f"record/{index}/{key}": value for key, value in row.items()})
    missing = [key for key, value in values.items() if normalized_text(value) not in text]
    return {
        "instrument": "PyMuPDF, unscored presence only",
        "pages": len(pages),
        "missing_surface_keys": missing,
        "limitation": "Presence is not evidence of row/field association or unclipped rendering",
    }


def prepare_references(
    corpus: Path, output: Path, document_ids: tuple[str, ...] = POPULATION
) -> Path:
    """Validate sources first and create a new unqualified snapshot without changing inputs."""
    if not document_ids or len(set(document_ids)) != len(document_ids):
        raise ValueError("Reference population must be nonempty and unique")
    if any(not re.fullmatch(r"invoice_\d{3}", doc_id) for doc_id in document_ids):
        raise ValueError("Invalid invoice identity")
    prepared = []
    for doc_id in document_ids:
        source = corpus / "docx" / f"{doc_id}.docx"
        pdf = corpus / "pdfs" / f"{doc_id}.pdf"
        metadata_path = corpus / "metadata" / f"{doc_id}.yaml"
        for path in [source, pdf, metadata_path]:
            if not path.is_file():
                raise ValueError(f"Reference gap: {doc_id}, missing {path.name}")
        frozen = {path: path.read_bytes() for path in [source, pdf, metadata_path]}
        reference = extract_reference(frozen[source])
        metadata = yaml.safe_load(frozen[metadata_path])
        cross_checks = {
            "invoice_number": metadata.get("invoice_number")
            == reference["fields"]["invoice_number"],
            "vendor": (metadata.get("vendor") or {}).get("name") == reference["fields"]["vendor"],
            "due_date": str(metadata.get("due_date")) == reference["fields"]["due_date"],
            "line_item_count": metadata.get("line_item_count") == len(reference["records"]),
            "grand_total": (
                Decimal(str(metadata["total"])) == Decimal(reference["fields"]["grand_total"])
                if metadata.get("total") is not None
                else False
            ),
        }
        presence = mechanical_presence(frozen[pdf], reference)
        prepared.append(
            (
                doc_id,
                reference,
                source,
                pdf,
                metadata_path,
                cross_checks,
                frozen,
                metadata,
                presence,
            )
        )
    output.mkdir(parents=True, exist_ok=True)
    identifier = uuid.uuid4().hex
    staging = output / f".incomplete-{identifier}"
    final = output / f"draft-{identifier}"
    staging.mkdir()
    protocol = reference_protocol()
    write_json_atomic(staging / "protocol.json", protocol)
    protocol_hash = digest(staging / "protocol.json")
    entries, files = [], {"protocol.json": protocol_hash}
    for (
        doc_id,
        reference,
        source,
        pdf,
        metadata_path,
        cross_checks,
        frozen,
        metadata,
        presence,
    ) in prepared:
        reference["document_id"] = doc_id
        for path, directory in [(source, "sources"), (pdf, "inputs"), (metadata_path, "metadata")]:
            target = staging / directory / path.name
            target.parent.mkdir(exist_ok=True)
            target.write_bytes(frozen[path])
            files[target.relative_to(staging).as_posix()] = digest(target)
        ref_path = staging / "references" / f"{doc_id}.json"
        write_json_atomic(ref_path, reference)
        files[ref_path.relative_to(staging).as_posix()] = digest(ref_path)
        entries.append(
            {
                "document_id": doc_id,
                "source_sha256": hashlib.sha256(frozen[source]).hexdigest(),
                "pdf_sha256": hashlib.sha256(frozen[pdf]).hexdigest(),
                "reference_sha256": digest(ref_path),
                "protocol_sha256": protocol_hash,
                "metadata_sha256": hashlib.sha256(frozen[metadata_path]).hexdigest(),
                "metadata_cross_checks": cross_checks,
                "template_style": metadata.get("template_style"),
                "complexity": metadata.get("complexity"),
                "mechanical": presence,
                "eligibility": "unqualified",
            }
        )
    manifest = {
        "status": "unqualified",
        "protocol_sha256": protocol_hash,
        "document_ids": list(document_ids),
        "documents": entries,
        "files": files,
    }
    write_json_atomic(staging / "manifest.json", manifest)
    for _, _, source, pdf, metadata_path, _, _, _, _ in prepared:
        for path, directory in [(source, "sources"), (pdf, "inputs"), (metadata_path, "metadata")]:
            if digest(path) != files[f"{directory}/{path.name}"]:
                raise ValueError("Source changed during reference preparation")
    validate_reference_bundle(staging)
    staging.rename(final)
    write_json_atomic(
        output / "latest-draft.json",
        {"bundle": final.name, "manifest_sha256": digest(final / "manifest.json")},
    )
    return final


def validate_reference_bundle(bundle: Path, require_qualified: bool = False) -> dict:
    manifest: dict = json.loads((bundle / "manifest.json").read_bytes())
    if manifest["status"] not in {"unqualified", "qualified"}:
        raise ValueError("Unknown reference bundle status")
    if require_qualified and manifest["status"] != "qualified":
        raise ValueError("References have not been qualified")
    ids = manifest["document_ids"]
    rows = manifest["documents"]
    if len(set(ids)) != len(ids) or sorted(ids) != sorted(row["document_id"] for row in rows):
        raise ValueError("Reference population has duplicate, missing or foreign identities")
    files = manifest["files"]
    for name, expected in files.items():
        path = (bundle / name).resolve()
        if not path.is_relative_to(bundle.resolve()) or digest(path) != expected:
            raise ValueError(f"Reference artifact hash/location mismatch: {name}")
    if files.get("protocol.json") != manifest["protocol_sha256"]:
        raise ValueError("Reference protocol identity mismatch")
    for row in rows:
        doc_id = row["document_id"]
        bindings = {
            f"sources/{doc_id}.docx": row["source_sha256"],
            f"inputs/{doc_id}.pdf": row["pdf_sha256"],
            f"references/{doc_id}.json": row["reference_sha256"],
            f"metadata/{doc_id}.yaml": row["metadata_sha256"],
            "protocol.json": row["protocol_sha256"],
        }
        if any(files.get(name) != value for name, value in bindings.items()):
            raise ValueError("Reference/metadata/rendition membership mismatch")
        reference = json.loads((bundle / f"references/{doc_id}.json").read_bytes())
        if reference["document_id"] != doc_id:
            raise ValueError("Reference facts belong to a different document")
        if require_qualified and row["eligibility"] != "eligible":
            raise ValueError("Unresolved truth in a qualified reference bundle")
    if require_qualified:
        if "qualification.json" not in files:
            raise ValueError("Qualified bundle lacks inspection evidence")
        decisions = json.loads((bundle / "qualification.json").read_bytes())["documents"]
        if sorted(ids) != sorted(d["document_id"] for d in decisions):
            raise ValueError("Qualification population mismatch")
        for row, decision in zip(rows, decisions, strict=True):
            if any(
                row[key] != decision[key]
                for key in (
                    "document_id",
                    "source_sha256",
                    "pdf_sha256",
                    "reference_sha256",
                    "protocol_sha256",
                )
            ):
                raise ValueError("Qualification belongs to another input/reference/protocol")
            if decision["verdict"] != "eligible":
                raise ValueError("Qualified bundle contains an unresolved decision")
    return manifest


def render_reference_pages(bundle: Path, output: Path) -> Path:
    """Render every fixed source PDF for an unscored, explicitly recorded page review."""
    manifest = validate_reference_bundle(bundle)
    if output.exists():
        raise FileExistsError("Render into a new location; retain prior inspection artifacts")
    output.mkdir(parents=True)
    artifacts = []
    for row in manifest["documents"]:
        doc_id = row["document_id"]
        pages = []
        with pymupdf.open(bundle / f"inputs/{doc_id}.pdf") as document:
            for number, page in enumerate(document, 1):
                image = page.get_pixmap(matrix=pymupdf.Matrix(1.25, 1.25), alpha=False).tobytes(
                    "png"
                )
                path = output / f"{doc_id}-page-{number}.png"
                path.write_bytes(image)
                pages.append({"page": number, "image": path.name, "sha256": digest(path)})
        artifacts.append(
            {
                **{
                    key: row[key]
                    for key in (
                        "document_id",
                        "source_sha256",
                        "pdf_sha256",
                        "reference_sha256",
                        "protocol_sha256",
                    )
                },
                "pages": pages,
            }
        )
    record = {
        "bundle_manifest_sha256": digest(bundle / "manifest.json"),
        "pymupdf_version": pymupdf.VersionBind,
        "scale": 1.25,
        "render_function_sha256": hashlib.sha256(
            inspect.getsource(render_reference_pages).encode()
        ).hexdigest(),
        "documents": artifacts,
    }
    write_json_atomic(output / "render-artifacts.json", record)
    return output / "render-artifacts.json"


def qualify_references(bundle: Path, rendered: Path, review_file: Path, output: Path) -> Path:
    """Require dated page judgments bound to the actual inputs before publishing eligibility."""
    manifest = validate_reference_bundle(bundle)
    if json.loads((bundle / "protocol.json").read_bytes()) != reference_protocol():
        raise ValueError("Reference protocol changed; prepare and inspect the new version")
    render = json.loads(rendered.read_bytes())
    review = json.loads(review_file.read_bytes())
    if render["bundle_manifest_sha256"] != digest(bundle / "manifest.json"):
        raise ValueError("Rendered review uses a different reference snapshot")
    if render["pymupdf_version"] != pymupdf.VersionBind or render["scale"] != 1.25:
        raise ValueError("Unrecognized render protocol")
    expected_renderer = hashlib.sha256(
        inspect.getsource(render_reference_pages).encode()
    ).hexdigest()
    if render["render_function_sha256"] != expected_renderer:
        raise ValueError("Rendering implementation changed")
    render_by_id = {d["document_id"]: d for d in render["documents"]}
    review_by_id = {d["document_id"]: d for d in review["documents"]}
    ids = set(manifest["document_ids"])
    if (
        set(render_by_id) != ids
        or set(review_by_id) != ids
        or len(render_by_id) != len(render["documents"])
        or len(review_by_id) != len(review["documents"])
    ):
        raise ValueError("Every declared document needs exactly one page review")
    decisions = []
    for row in manifest["documents"]:
        doc_id = row["document_id"]
        rendered_doc, decision = render_by_id[doc_id], review_by_id[doc_id]
        for key in ("source_sha256", "pdf_sha256", "reference_sha256", "protocol_sha256"):
            if row[key] != rendered_doc[key] or row[key] != decision[key]:
                raise ValueError("Stale source/rendition/reference/protocol qualification")
        with pymupdf.open(bundle / f"inputs/{doc_id}.pdf") as document:
            expected_pages = list(range(1, len(document) + 1))
            if (
                decision["reviewed_pages"] != expected_pages
                or [p["page"] for p in rendered_doc["pages"]] != expected_pages
            ):
                raise ValueError("Every physical page requires review")
            for page, image in zip(document, rendered_doc["pages"], strict=True):
                data = page.get_pixmap(matrix=pymupdf.Matrix(1.25, 1.25), alpha=False).tobytes(
                    "png"
                )
                if hashlib.sha256(data).hexdigest() != image["sha256"]:
                    raise ValueError("Reviewed rendering does not match frozen PDF")
                if digest(rendered.parent / image["image"]) != image["sha256"]:
                    raise ValueError("Inspection image changed after review")
        if (
            decision["verdict"] != "eligible"
            or not decision.get("associations_checked")
            or not decision.get("layout_checked")
            or not decision.get("notes")
        ):
            raise ValueError("Unresolved source/PDF truth or layout blocks strong qualification")
        reference = json.loads((bundle / f"references/{doc_id}.json").read_bytes())
        if decision.get("observed_record_count") != len(reference["records"]):
            raise ValueError("Rendered record count disagrees with source")
        if decision.get("observed_headers") != reference["source_headers"]:
            raise ValueError("Rendered header roles disagree with source")
        fresh_presence = mechanical_presence(bundle / f"inputs/{doc_id}.pdf", reference)
        if fresh_presence != row["mechanical"]:
            raise ValueError("Reference presence diagnostic changed")
        resolutions = decision.get("mechanical_resolutions", {})
        for key in row["mechanical"]["missing_surface_keys"]:
            if key not in resolutions or not resolutions[key].get("explanation"):
                raise ValueError("A text-presence miss needs explicit source-page resolution")
            if not set(resolutions[key]["pages"]).issubset(decision["reviewed_pages"]):
                raise ValueError("Presence resolution points outside reviewed pages")
        for key, agrees in row["metadata_cross_checks"].items():
            if not agrees and not decision.get("metadata_resolutions", {}).get(key):
                raise ValueError("Metadata disagreement needs an explicit source/PDF decision")
        decisions.append({**decision, "rendered_pages": rendered_doc["pages"]})
    identifier = uuid.uuid4().hex
    output.mkdir(parents=True, exist_ok=True)
    staging, final = output / f".incomplete-{identifier}", output / f"qualified-{identifier}"
    shutil.copytree(bundle, staging)
    qualification = {
        "reviewer": review["reviewer"],
        "reviewed_at": review["reviewed_at"],
        "method": review["method"],
        "documents": decisions,
    }
    write_json_atomic(staging / "qualification.json", qualification)
    manifest["files"]["qualification.json"] = digest(staging / "qualification.json")
    manifest["status"] = "qualified"
    for row in manifest["documents"]:
        row["eligibility"] = "eligible"
    write_json_atomic(staging / "manifest.json", manifest)
    validate_reference_bundle(staging, require_qualified=True)
    staging.rename(final)
    write_json_atomic(
        output / "latest-qualified.json",
        {"bundle": final.name, "manifest_sha256": digest(final / "manifest.json")},
    )
    return final


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("action", choices=["prepare", "render", "qualify"])
    parser.add_argument("--corpus", type=Path)
    parser.add_argument("--bundle", type=Path)
    parser.add_argument("--rendered", type=Path)
    parser.add_argument("--review", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if args.action == "prepare":
        if args.corpus is None:
            parser.error("prepare requires --corpus")
        print(prepare_references(args.corpus, args.output))
    elif args.action == "render":
        if args.bundle is None:
            parser.error("render requires --bundle")
        print(render_reference_pages(args.bundle, args.output))
    else:
        if None in (args.bundle, args.rendered, args.review):
            parser.error("qualify requires --bundle, --rendered and --review")
        print(qualify_references(args.bundle, args.rendered, args.review, args.output))


if __name__ == "__main__":
    main()
