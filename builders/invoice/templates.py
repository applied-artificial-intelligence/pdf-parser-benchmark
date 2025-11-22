"""DOCX invoice template generation.

This module creates three professional invoice templates:
1. Enterprise Template - Clean, corporate style
2. Small Business Template - Friendly, accessible style
3. Economy Template - Minimal, cost-conscious style
"""

from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from builders.invoice.data_generators import CompanyInfo, LineItem
from builders.invoice.tax_calculators import TaxResult


def add_table_borders(table):
    """Add borders to all cells in a table."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)

    tblBorders = OxmlElement("w:tblBorders")
    for border_name in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        border = OxmlElement(f"w:{border_name}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "000000")
        tblBorders.append(border)

    tblPr.append(tblBorders)


class EnterpriseTemplate:
    """Clean, corporate-style invoice template."""

    def __init__(self):
        self.doc = Document()
        # Set narrow margins
        sections = self.doc.sections
        for section in sections:
            section.top_margin = Inches(0.5)
            section.bottom_margin = Inches(0.5)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)

    def create_invoice(
        self,
        invoice_number: str,
        invoice_date: datetime,
        due_date: datetime,
        vendor: CompanyInfo,
        customer: CompanyInfo,
        line_items: list[LineItem],
        tax_result: TaxResult,
        payment_terms: str,
        notes: str = "",
        po_number: str = "",
    ) -> Document:
        """Generate enterprise-style invoice."""

        # Header - INVOICE title
        title = self.doc.add_paragraph()
        title_run = title.add_run("INVOICE")
        title_run.font.size = Pt(28)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(0, 51, 102)  # Dark blue
        title.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        self.doc.add_paragraph()  # Spacing

        # Company info header table
        header_table = self.doc.add_table(rows=1, cols=2)
        header_table.autofit = False
        header_table.allow_autofit = False

        # Left column - Vendor (From)
        left_cell = header_table.rows[0].cells[0]
        left_para = left_cell.paragraphs[0]

        vendor_run = left_para.add_run(f"{vendor.name}\n")
        vendor_run.font.bold = True
        vendor_run.font.size = Pt(11)

        left_para.add_run(f"{vendor.address}\n")
        left_para.add_run(f"{vendor.city}, {vendor.state_province} {vendor.postal_code}\n")
        if vendor.country:
            left_para.add_run(f"{vendor.country}\n")
        left_para.add_run("\n")
        left_para.add_run(f"Phone: {vendor.phone}\n")
        left_para.add_run(f"Email: {vendor.email}\n")
        if vendor.tax_id:
            left_para.add_run(f"{vendor.tax_id}\n")

        # Right column - Invoice details
        right_cell = header_table.rows[0].cells[1]
        right_para = right_cell.paragraphs[0]
        right_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        right_para.add_run("Invoice #: ").font.bold = True
        right_para.add_run(f"{invoice_number}\n")

        right_para.add_run("Date: ").font.bold = True
        right_para.add_run(f"{invoice_date.strftime('%B %d, %Y')}\n")

        right_para.add_run("Due Date: ").font.bold = True
        right_para.add_run(f"{due_date.strftime('%B %d, %Y')}\n")

        if po_number:
            right_para.add_run("\nPO Number: ").font.bold = True
            right_para.add_run(f"{po_number}\n")

        self.doc.add_paragraph()  # Spacing

        # Bill To section
        bill_to_para = self.doc.add_paragraph()
        bill_to_run = bill_to_para.add_run("BILL TO")
        bill_to_run.font.bold = True
        bill_to_run.font.size = Pt(10)
        bill_to_run.font.color.rgb = RGBColor(0, 51, 102)

        customer_para = self.doc.add_paragraph()
        customer_name_run = customer_para.add_run(f"{customer.name}\n")
        customer_name_run.font.bold = True
        customer_para.add_run(f"{customer.address}\n")
        customer_para.add_run(
            f"{customer.city}, {customer.state_province} {customer.postal_code}\n"
        )
        if customer.country:
            customer_para.add_run(f"{customer.country}\n")

        self.doc.add_paragraph()  # Spacing

        # Line items table
        items_table = self.doc.add_table(rows=1, cols=5)
        items_table.style = "Light Grid Accent 1"

        # Header row
        header_cells = items_table.rows[0].cells
        headers = ["Description", "Quantity", "Unit", "Unit Price", "Amount"]
        for i, header_text in enumerate(headers):
            cell_para = header_cells[i].paragraphs[0]
            cell_run = cell_para.add_run(header_text)
            cell_run.font.bold = True
            cell_run.font.size = Pt(10)
            if i >= 3:  # Right-align price columns
                cell_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        # Line items
        for item in line_items:
            row_cells = items_table.add_row().cells
            row_cells[0].text = item.description
            row_cells[1].text = str(item.quantity)
            row_cells[2].text = item.unit

            # Format currency
            price_para = row_cells[3].paragraphs[0]
            price_para.text = f"${item.unit_price:,.2f}"
            price_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

            amount_para = row_cells[4].paragraphs[0]
            amount_para.text = f"${item.amount:,.2f}"
            amount_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        add_table_borders(items_table)

        self.doc.add_paragraph()  # Spacing

        # Totals table (right-aligned)
        totals_table = self.doc.add_table(rows=3, cols=2)
        totals_table.autofit = False
        totals_table.allow_autofit = False

        # Set column widths
        for row in totals_table.rows:
            row.cells[0].width = Inches(1.5)
            row.cells[1].width = Inches(1.5)

        # Move table to right
        totals_table.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        # Subtotal
        totals_table.rows[0].cells[0].text = "Subtotal:"
        totals_table.rows[0].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        totals_table.rows[0].cells[1].text = f"${tax_result.subtotal:,.2f}"
        totals_table.rows[0].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

        # Tax
        totals_table.rows[1].cells[0].text = f"{tax_result.tax_label}:"
        totals_table.rows[1].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        totals_table.rows[1].cells[1].text = f"${tax_result.tax_amount:,.2f}"
        totals_table.rows[1].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

        # Total
        total_label_para = totals_table.rows[2].cells[0].paragraphs[0]
        total_label_run = total_label_para.add_run("TOTAL:")
        total_label_run.font.bold = True
        total_label_run.font.size = Pt(12)
        total_label_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        total_amount_para = totals_table.rows[2].cells[1].paragraphs[0]
        total_amount_run = total_amount_para.add_run(f"${tax_result.total:,.2f}")
        total_amount_run.font.bold = True
        total_amount_run.font.size = Pt(12)
        total_amount_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        self.doc.add_paragraph()  # Spacing

        # Payment terms
        terms_para = self.doc.add_paragraph()
        terms_label = terms_para.add_run("Payment Terms: ")
        terms_label.font.bold = True
        terms_para.add_run(payment_terms)

        # Tax number if available
        if tax_result.tax_number:
            tax_num_para = self.doc.add_paragraph()
            tax_num_para.add_run(tax_result.tax_number)
            tax_num_para.style = "Intense Quote"

        # Notes
        if notes:
            self.doc.add_paragraph()
            notes_para = self.doc.add_paragraph()
            notes_label = notes_para.add_run("Notes:\n")
            notes_label.font.bold = True
            notes_para.add_run(notes)

        return self.doc


class SmallBusinessTemplate:
    """Friendly, accessible invoice template for small businesses."""

    def __init__(self):
        self.doc = Document()
        sections = self.doc.sections
        for section in sections:
            section.top_margin = Inches(0.75)
            section.bottom_margin = Inches(0.75)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)

    def create_invoice(
        self,
        invoice_number: str,
        invoice_date: datetime,
        due_date: datetime,
        vendor: CompanyInfo,
        customer: CompanyInfo,
        line_items: list[LineItem],
        tax_result: TaxResult,
        payment_terms: str,
        notes: str = "",
        po_number: str = "",
    ) -> Document:
        """Generate small business-style invoice."""

        # Friendly header
        title = self.doc.add_paragraph()
        title_run = title.add_run("Invoice")
        title_run.font.size = Pt(24)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(51, 102, 153)  # Friendly blue
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Vendor info (centered)
        vendor_para = self.doc.add_paragraph()
        vendor_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        vendor_name_run = vendor_para.add_run(f"{vendor.name}\n")
        vendor_name_run.font.size = Pt(12)
        vendor_name_run.font.bold = True
        vendor_para.add_run(f"{vendor.address}\n")
        vendor_para.add_run(f"{vendor.city}, {vendor.state_province} {vendor.postal_code}\n")
        vendor_para.add_run(f"{vendor.phone} | {vendor.email}\n")

        self.doc.add_paragraph()  # Spacing

        # Invoice details (centered box)
        details_table = self.doc.add_table(rows=4, cols=2)
        details_table.alignment = WD_ALIGN_PARAGRAPH.CENTER

        details_table.rows[0].cells[0].text = "Invoice Number:"
        details_table.rows[0].cells[1].text = invoice_number

        details_table.rows[1].cells[0].text = "Invoice Date:"
        details_table.rows[1].cells[1].text = invoice_date.strftime("%B %d, %Y")

        details_table.rows[2].cells[0].text = "Payment Due:"
        details_table.rows[2].cells[1].text = due_date.strftime("%B %d, %Y")

        details_table.rows[3].cells[0].text = "Payment Terms:"
        details_table.rows[3].cells[1].text = payment_terms

        add_table_borders(details_table)

        self.doc.add_paragraph()  # Spacing

        # Bill To
        bill_to_para = self.doc.add_paragraph()
        bill_to_run = bill_to_para.add_run("Bill To:")
        bill_to_run.font.bold = True
        bill_to_run.font.size = Pt(11)

        customer_para = self.doc.add_paragraph()
        customer_name_run = customer_para.add_run(f"{customer.name}\n")
        customer_name_run.font.bold = True
        customer_para.add_run(f"{customer.address}\n")
        customer_para.add_run(
            f"{customer.city}, {customer.state_province} {customer.postal_code}\n"
        )

        self.doc.add_paragraph()  # Spacing

        # Line items
        items_table = self.doc.add_table(rows=1, cols=4)

        # Header
        header_cells = items_table.rows[0].cells
        headers = ["Description", "Qty", "Rate", "Amount"]
        for i, header_text in enumerate(headers):
            cell_para = header_cells[i].paragraphs[0]
            cell_run = cell_para.add_run(header_text)
            cell_run.font.bold = True
            if i >= 2:
                cell_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        # Items
        for item in line_items:
            row_cells = items_table.add_row().cells
            row_cells[0].text = f"{item.description} ({item.unit})"
            row_cells[1].text = str(item.quantity)

            rate_para = row_cells[2].paragraphs[0]
            rate_para.text = f"${item.unit_price:,.2f}"
            rate_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

            amount_para = row_cells[3].paragraphs[0]
            amount_para.text = f"${item.amount:,.2f}"
            amount_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        add_table_borders(items_table)

        self.doc.add_paragraph()  # Spacing

        # Totals
        totals_table = self.doc.add_table(rows=3, cols=2)
        totals_table.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        totals_table.rows[0].cells[0].text = "Subtotal:"
        totals_table.rows[0].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        totals_table.rows[0].cells[1].text = f"${tax_result.subtotal:,.2f}"
        totals_table.rows[0].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

        totals_table.rows[1].cells[0].text = f"{tax_result.tax_label}:"
        totals_table.rows[1].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        totals_table.rows[1].cells[1].text = f"${tax_result.tax_amount:,.2f}"
        totals_table.rows[1].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

        total_para_label = totals_table.rows[2].cells[0].paragraphs[0]
        total_run_label = total_para_label.add_run("Total:")
        total_run_label.font.bold = True
        total_para_label.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        total_para_amount = totals_table.rows[2].cells[1].paragraphs[0]
        total_run_amount = total_para_amount.add_run(f"${tax_result.total:,.2f}")
        total_run_amount.font.bold = True
        total_para_amount.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        # Notes
        if notes:
            self.doc.add_paragraph()
            notes_para = self.doc.add_paragraph()
            notes_para.add_run(notes)
            notes_para.style = "Intense Quote"

        return self.doc


class EconomyTemplate:
    """Minimal, cost-conscious invoice template."""

    def __init__(self):
        self.doc = Document()
        sections = self.doc.sections
        for section in sections:
            section.top_margin = Inches(0.5)
            section.bottom_margin = Inches(0.5)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)

    def create_invoice(
        self,
        invoice_number: str,
        invoice_date: datetime,
        due_date: datetime,
        vendor: CompanyInfo,
        customer: CompanyInfo,
        line_items: list[LineItem],
        tax_result: TaxResult,
        payment_terms: str,
        notes: str = "",
        po_number: str = "",
    ) -> Document:
        """Generate economy-style invoice."""

        # Simple header
        header = self.doc.add_paragraph()
        header_run = header.add_run("INVOICE")
        header_run.font.size = Pt(16)
        header_run.font.bold = True

        # Vendor and customer on same line
        info_table = self.doc.add_table(rows=1, cols=2)

        # Vendor (FROM)
        from_cell = info_table.rows[0].cells[0]
        from_para = from_cell.paragraphs[0]
        from_para.add_run("FROM:\n").font.bold = True
        from_para.add_run(f"{vendor.name}\n")
        from_para.add_run(f"{vendor.address}\n")
        from_para.add_run(f"{vendor.city}, {vendor.state_province} {vendor.postal_code}\n")
        from_para.add_run(f"{vendor.phone}\n")

        # Customer (TO)
        to_cell = info_table.rows[0].cells[1]
        to_para = to_cell.paragraphs[0]
        to_para.add_run("TO:\n").font.bold = True
        to_para.add_run(f"{customer.name}\n")
        to_para.add_run(f"{customer.address}\n")
        to_para.add_run(f"{customer.city}, {customer.state_province} {customer.postal_code}\n")

        self.doc.add_paragraph()  # Spacing

        # Invoice details (one line)
        details_para = self.doc.add_paragraph()
        details_para.add_run(f"Invoice #: {invoice_number} | ")
        details_para.add_run(f"Date: {invoice_date.strftime('%m/%d/%Y')} | ")
        details_para.add_run(f"Due: {due_date.strftime('%m/%d/%Y')} | ")
        details_para.add_run(f"Terms: {payment_terms}")

        self.doc.add_paragraph()  # Spacing

        # Line items (simple table)
        items_table = self.doc.add_table(rows=1, cols=3)

        header_cells = items_table.rows[0].cells
        header_cells[0].text = "Description"
        header_cells[1].text = "Qty"
        header_cells[2].text = "Amount"

        for cell in header_cells:
            cell.paragraphs[0].runs[0].font.bold = True

        for item in line_items:
            row_cells = items_table.add_row().cells
            row_cells[0].text = f"{item.description} @ ${item.unit_price}/{item.unit}"
            row_cells[1].text = str(item.quantity)
            row_cells[2].text = f"${item.amount:,.2f}"

        add_table_borders(items_table)

        self.doc.add_paragraph()  # Spacing

        # Totals (simple)
        totals_para = self.doc.add_paragraph()
        totals_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        totals_para.add_run(f"Subtotal: ${tax_result.subtotal:,.2f}\n")
        totals_para.add_run(f"{tax_result.tax_label}: ${tax_result.tax_amount:,.2f}\n")

        total_run = totals_para.add_run(f"TOTAL: ${tax_result.total:,.2f}")
        total_run.font.bold = True
        total_run.font.size = Pt(12)

        # Notes if any
        if notes:
            self.doc.add_paragraph()
            self.doc.add_paragraph(notes)

        return self.doc


def get_template(style: str):
    """Get invoice template by style.

    Args:
        style: One of "enterprise", "small_business", "economy"

    Returns:
        Template instance
    """
    templates = {
        "enterprise": EnterpriseTemplate,
        "small_business": SmallBusinessTemplate,
        "economy": EconomyTemplate,
    }
    template_class = templates.get(style, EnterpriseTemplate)
    return template_class()
