"""Legal contract templates using python-docx for authentic formatting.

Creates properly formatted legal contracts with:
- Title page with contract name and parties
- Recitals (WHEREAS clauses)
- Numbered sections using legal numbering system
- Signature blocks
- Professional formatting (fonts, spacing, alignment)

Supports 5 contract types:
- SaaS Agreement
- Employment Agreement
- Non-Disclosure Agreement
- Consulting Agreement
- Service Agreement
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from typing import Dict, List, Optional
import random
from datetime import datetime

from .contract_config import (
    SAAS_CONFIG, EMPLOYMENT_CONFIG, NDA_CONFIG,
    CONSULTING_CONFIG, SERVICE_CONFIG, ContractConfig
)
from .clause_library import ClauseLibrary, DefinedTermsTracker, ClauseVariant
from .numbering import ContractNumbering


class BaseContractTemplate:
    """Base class for all contract templates."""

    def __init__(self, contract_config: ContractConfig):
        self.config = contract_config
        self.doc = Document()
        self.clause_lib = ClauseLibrary()
        self.numbering = ContractNumbering()
        self.terms_tracker = DefinedTermsTracker()
        self.variables: Dict[str, any] = {}

        # Set up document formatting
        self._setup_document_styles()

    def _setup_document_styles(self):
        """Configure document-wide styles."""
        # Set up default font
        style = self.doc.styles['Normal']
        font = style.font
        font.name = 'Times New Roman'
        font.size = Pt(12)

        # Set up heading styles
        for i in range(1, 4):
            heading_style = self.doc.styles[f'Heading {i}']
            heading_font = heading_style.font
            heading_font.name = 'Times New Roman'
            heading_font.bold = True
            if i == 1:
                heading_font.size = Pt(14)
            else:
                heading_font.size = Pt(12)

    def _add_title_page(self, title: str, parties: List[Dict[str, str]], effective_date: str):
        """Add contract title page.

        Args:
            title: Contract title (e.g., "SOFTWARE AS A SERVICE AGREEMENT")
            parties: List of party dicts with keys: name, type, address
            effective_date: Effective date string
        """
        # Title
        title_para = self.doc.add_paragraph()
        title_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        title_run = title_para.add_run(title.upper())
        title_run.bold = True
        title_run.font.size = Pt(16)

        self.doc.add_paragraph()  # Spacing

        # Effective date
        date_para = self.doc.add_paragraph()
        date_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        date_para.add_run(f"Effective Date: {effective_date}")

        self.doc.add_paragraph()  # Spacing

        # Parties
        parties_para = self.doc.add_paragraph()
        parties_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        parties_run = parties_para.add_run("BY AND BETWEEN")
        parties_run.bold = True

        self.doc.add_paragraph()  # Spacing

        for i, party in enumerate(parties):
            party_para = self.doc.add_paragraph()
            party_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            party_para.add_run(f'{party["name"]}\n')
            if "address" in party:
                party_para.add_run(f'{party["address"]}\n')
            party_type_run = party_para.add_run(f'("{party["type"]}")')
            party_type_run.italic = True

            if i < len(parties) - 1:
                and_para = self.doc.add_paragraph()
                and_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                and_para.add_run("and")
                self.doc.add_paragraph()  # Spacing

        # Page break
        self.doc.add_page_break()

    def _add_recitals(self, recitals: List[str]):
        """Add WHEREAS recital clauses.

        Args:
            recitals: List of recital texts
        """
        recitals_heading = self.doc.add_paragraph()
        recitals_run = recitals_heading.add_run("RECITALS")
        recitals_run.bold = True
        recitals_run.font.size = Pt(12)
        recitals_heading.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        self.doc.add_paragraph()  # Spacing

        for recital in recitals:
            para = self.doc.add_paragraph()
            para.add_run(f"WHEREAS, {recital};")
            para.paragraph_format.left_indent = Inches(0.5)

        # NOW THEREFORE clause
        self.doc.add_paragraph()
        now_para = self.doc.add_paragraph()
        now_para.add_run("NOW, THEREFORE, in consideration of the mutual covenants and agreements set forth herein and for other good and valuable consideration, the receipt and sufficiency of which are hereby acknowledged, the parties agree as follows:")
        now_para.paragraph_format.left_indent = Inches(0.5)

        self.doc.add_paragraph()  # Spacing

    def _add_numbered_section(self, level: int, title: str = "", content: str = ""):
        """Add a numbered section using legal numbering.

        Args:
            level: Numbering level (1-5)
            title: Section title
            content: Section content
        """
        item = self.numbering.add_section(level, title, content)

        para = self.doc.add_paragraph()

        # Format based on level
        if level == 1:
            # Article/Major section - bold title
            run = para.add_run(f"{item.number_str}. {title}")
            run.bold = True
            run.font.size = Pt(12)
            para.paragraph_format.space_before = Pt(12)
            para.paragraph_format.space_after = Pt(6)
        elif level == 2:
            # Subsection - bold title
            run = para.add_run(f"{item.number_str} {title}")
            run.bold = True
            para.paragraph_format.left_indent = Inches(0.25)
            para.paragraph_format.space_before = Pt(6)
            para.paragraph_format.space_after = Pt(3)
        elif level == 3:
            # Paragraph
            para.paragraph_format.left_indent = Inches(0.5)
            para.add_run(f"{item.number_str} {content}")
        elif level == 4:
            # Subparagraph with letter
            para.paragraph_format.left_indent = Inches(0.75)
            para.add_run(f"{item.number_str} {content}")
        else:  # level 5
            # Clause with Roman numeral
            para.paragraph_format.left_indent = Inches(1.0)
            para.add_run(f"{item.number_str} {content}")

    def _add_clause_with_substitution(
        self,
        section_name: str,
        level: int = 3,
        variant: ClauseVariant = ClauseVariant.STANDARD,
        contract_type: Optional[str] = None,
    ):
        """Add a clause from the library with variable substitution.

        Args:
            section_name: Section name to get clause for
            level: Numbering level for this clause
            variant: Clause variant (strict, standard, permissive)
            contract_type: Contract type filter
        """
        clause = self.clause_lib.get_clause(section_name, contract_type, variant)
        if not clause:
            # print(f"Warning: No clause found for section '{section_name}'")
            return

        # Substitute variables
        content = self.clause_lib.substitute_variables(clause.text, self.variables)

        # Track defined terms
        if clause.defined_terms:
            self.terms_tracker.add_terms(clause.defined_terms)

        # Add to document
        self._add_numbered_section(level, content=content)

    def _add_signature_block(self, parties: List[Dict[str, str]]):
        """Add signature blocks for all parties.

        Args:
            parties: List of party dicts with keys: name, type
        """
        self.doc.add_paragraph()  # Spacing
        self.doc.add_paragraph()

        sig_para = self.doc.add_paragraph()
        sig_run = sig_para.add_run("IN WITNESS WHEREOF, the parties have executed this Agreement as of the Effective Date.")
        sig_para.paragraph_format.space_before = Pt(12)

        self.doc.add_paragraph()
        self.doc.add_paragraph()

        for party in parties:
            # Party name (all caps)
            party_para = self.doc.add_paragraph()
            party_run = party_para.add_run(party["name"].upper())
            party_run.bold = True

            self.doc.add_paragraph()  # Spacing

            # Signature line
            sig_line = self.doc.add_paragraph()
            sig_line.add_run("By: _" + "_" * 50)

            # Name/Title
            name_para = self.doc.add_paragraph()
            name_para.add_run("Name: _" + "_" * 45)

            title_para = self.doc.add_paragraph()
            title_para.add_run("Title: _" + "_" * 45)

            date_para = self.doc.add_paragraph()
            date_para.add_run("Date: _" + "_" * 45)

            self.doc.add_paragraph()  # Spacing between parties
            self.doc.add_paragraph()

    def generate(self, variables: Dict[str, any]) -> Document:
        """Generate the contract document.

        Args:
            variables: Dictionary of variables for substitution

        Returns:
            Generated Document object
        """
        raise NotImplementedError("Subclasses must implement generate()")


class SaaSAgreementTemplate(BaseContractTemplate):
    """SaaS Agreement template."""

    def __init__(self):
        super().__init__(SAAS_CONFIG)

    def generate(self, variables: Dict[str, any]) -> Document:
        """Generate SaaS Agreement.

        Required variables:
        - provider: Provider company name
        - customer: Customer company name
        - provider_address: Provider address
        - customer_address: Customer address
        - effective_date: Contract effective date
        - platform_name: Platform/product name
        - service_purpose: What the service does
        - uptime_pct: Uptime percentage (e.g., "99.9")
        - monthly_fee: Monthly subscription fee
        - jurisdiction: Governing law jurisdiction
        """
        self.variables = variables

        # Title page
        parties = [
            {"name": variables["provider"], "type": "Provider", "address": variables.get("provider_address", "")},
            {"name": variables["customer"], "type": "Customer", "address": variables.get("customer_address", "")},
        ]
        self._add_title_page("SOFTWARE AS A SERVICE AGREEMENT", parties, variables["effective_date"])

        # Recitals
        recitals = [
            f'{variables["provider"]} offers a cloud-based software platform',
            f'{variables["customer"]} desires to subscribe to and use the Services',
            "the parties wish to establish terms governing such use",
        ]
        self._add_recitals(recitals)

        # Articles
        self._add_numbered_section(1, "Definitions")
        self._add_clause_with_substitution("Service Description", level=3, contract_type="saas")

        self._add_numbered_section(1, "Services")
        self._add_numbered_section(2, "Service Level Agreement")
        self._add_clause_with_substitution("Service Level Agreement", level=3, contract_type="saas")
        self._add_numbered_section(2, "Support Services")
        self._add_clause_with_substitution("Support Services", level=3, contract_type="saas")

        self._add_numbered_section(1, "Subscription and Fees")
        self._add_numbered_section(2, "Subscription Plans")
        self._add_clause_with_substitution("Subscription Plans", level=3, contract_type="saas")
        self._add_numbered_section(2, "Payment Terms")
        self._add_clause_with_substitution("Payment Terms", level=3, contract_type="saas")

        self._add_numbered_section(1, "Data and Privacy")
        self._add_numbered_section(2, "Data Ownership")
        self._add_clause_with_substitution("Data Ownership", level=3, contract_type="saas")
        self._add_numbered_section(2, "Data Security")
        self._add_clause_with_substitution("Data Security", level=3, contract_type="saas")

        self._add_numbered_section(1, "Term and Termination")
        self._add_numbered_section(2, "Initial Term")
        self._add_clause_with_substitution("Term and Renewal", level=3, contract_type="service")

        self._add_numbered_section(1, "Limitation of Liability")
        self._add_clause_with_substitution("Limitation of Liability", level=3, contract_type="saas")

        self._add_numbered_section(1, "General Provisions")
        self._add_numbered_section(2, "Governing Law")
        self._add_clause_with_substitution("Governing Law", level=3)
        self._add_numbered_section(2, "Entire Agreement")
        self._add_clause_with_substitution("Entire Agreement", level=3)
        self._add_numbered_section(2, "Notices")
        self._add_clause_with_substitution("Notices", level=3)

        # Signature block
        self._add_signature_block(parties)

        return self.doc


class EmploymentAgreementTemplate(BaseContractTemplate):
    """Employment Agreement template."""

    def __init__(self):
        super().__init__(EMPLOYMENT_CONFIG)

    def generate(self, variables: Dict[str, any]) -> Document:
        """Generate Employment Agreement.

        Required variables:
        - company: Company name
        - employee: Employee name
        - position_title: Job title
        - annual_salary: Annual salary amount
        - jurisdiction: Governing law jurisdiction
        - effective_date: Start date
        """
        self.variables = variables

        parties = [
            {"name": variables["company"], "type": "Company"},
            {"name": variables["employee"], "type": "Employee"},
        ]
        self._add_title_page("EMPLOYMENT AGREEMENT", parties, variables["effective_date"])

        recitals = [
            f'{variables["company"]} desires to employ {variables["employee"]} in the specified position',
            f'{variables["employee"]} wishes to accept such employment',
            "the parties desire to set forth the terms of the employment relationship",
        ]
        self._add_recitals(recitals)

        self._add_numbered_section(1, "Position and Duties")
        self._add_clause_with_substitution("Position and Duties", level=3, contract_type="employment")

        self._add_numbered_section(1, "Compensation")
        self._add_numbered_section(2, "Base Salary")
        self._add_clause_with_substitution("Base Salary", level=3, contract_type="employment")

        self._add_numbered_section(1, "At-Will Employment")
        self._add_clause_with_substitution("At-Will Employment", level=3, contract_type="employment")

        self._add_numbered_section(1, "Confidentiality")
        self._add_clause_with_substitution("Confidentiality", level=3, contract_type="employment")

        self._add_numbered_section(1, "Intellectual Property Assignment")
        self._add_clause_with_substitution("Intellectual Property Assignment", level=3, contract_type="employment")

        self._add_numbered_section(1, "General Provisions")
        self._add_numbered_section(2, "Governing Law")
        self._add_clause_with_substitution("Governing Law", level=3)
        self._add_numbered_section(2, "Entire Agreement")
        self._add_clause_with_substitution("Entire Agreement", level=3)

        self._add_signature_block(parties)

        return self.doc


class NDATemplate(BaseContractTemplate):
    """Non-Disclosure Agreement template."""

    def __init__(self):
        super().__init__(NDA_CONFIG)

    def generate(self, variables: Dict[str, any]) -> Document:
        """Generate NDA.

        Required variables:
        - disclosing_party: Disclosing party name
        - receiving_party: Receiving party name
        - purpose: Purpose of disclosure
        - effective_date: Effective date
        - jurisdiction: Governing law jurisdiction
        """
        self.variables = variables

        parties = [
            {"name": variables["disclosing_party"], "type": "Disclosing Party"},
            {"name": variables["receiving_party"], "type": "Receiving Party"},
        ]
        self._add_title_page("NON-DISCLOSURE AGREEMENT", parties, variables["effective_date"])

        recitals = [
            "the parties contemplate entering into business discussions",
            "disclosures of Confidential Information may be necessary",
            "the parties desire to protect the confidentiality of such information",
        ]
        self._add_recitals(recitals)

        self._add_numbered_section(1, "Definitions")
        self._add_numbered_section(2, "Confidential Information")
        self._add_clause_with_substitution("Definition of Confidential Information", level=3, contract_type="nda")
        self._add_clause_with_substitution("Exclusions from Confidential Information", level=3, contract_type="nda")

        self._add_numbered_section(1, "Obligations")
        self._add_clause_with_substitution("Non-Disclosure", level=3, contract_type="nda")
        self._add_clause_with_substitution("Protection Measures", level=3, contract_type="nda")

        self._add_numbered_section(1, "Return of Materials")
        self._add_clause_with_substitution("Return of Materials", level=3, contract_type="nda")

        self._add_numbered_section(1, "Remedies")
        self._add_clause_with_substitution("Remedies", level=3, contract_type="nda")

        self._add_numbered_section(1, "General Provisions")
        self._add_numbered_section(2, "Governing Law")
        self._add_clause_with_substitution("Governing Law", level=3)
        self._add_numbered_section(2, "Entire Agreement")
        self._add_clause_with_substitution("Entire Agreement", level=3)

        self._add_signature_block(parties)

        return self.doc


class ConsultingAgreementTemplate(BaseContractTemplate):
    """Consulting Agreement template."""

    def __init__(self):
        super().__init__(CONSULTING_CONFIG)

    def generate(self, variables: Dict[str, any]) -> Document:
        """Generate Consulting Agreement.

        Required variables:
        - client: Client company name
        - consultant: Consultant name
        - hourly_rate: Hourly rate
        - effective_date: Start date
        - jurisdiction: Governing law jurisdiction
        """
        self.variables = variables

        parties = [
            {"name": variables["client"], "type": "Client"},
            {"name": variables["consultant"], "type": "Consultant"},
        ]
        self._add_title_page("CONSULTING AGREEMENT", parties, variables["effective_date"])

        recitals = [
            f'{variables["client"]} desires to engage {variables["consultant"]} to provide certain services',
            f'{variables["consultant"]} has expertise and is willing to provide such services',
            "the parties wish to establish terms governing the engagement",
        ]
        self._add_recitals(recitals)

        self._add_numbered_section(1, "Services")
        self._add_clause_with_substitution("Scope of Services", level=3, contract_type="consulting")
        self._add_clause_with_substitution("Deliverables", level=3, contract_type="consulting")

        self._add_numbered_section(1, "Compensation")
        self._add_clause_with_substitution("Fees", level=3, contract_type="consulting")
        self._add_clause_with_substitution("Expenses", level=3, contract_type="consulting")

        self._add_numbered_section(1, "Independent Contractor Relationship")
        self._add_clause_with_substitution("Independent Contractor Relationship", level=3, contract_type="consulting")

        self._add_numbered_section(1, "Intellectual Property")
        self._add_clause_with_substitution("Work Product Ownership", level=3, contract_type="consulting")

        self._add_numbered_section(1, "General Provisions")
        self._add_numbered_section(2, "Governing Law")
        self._add_clause_with_substitution("Governing Law", level=3)
        self._add_numbered_section(2, "Entire Agreement")
        self._add_clause_with_substitution("Entire Agreement", level=3)

        self._add_signature_block(parties)

        return self.doc


class ServiceAgreementTemplate(BaseContractTemplate):
    """Service Agreement template."""

    def __init__(self):
        super().__init__(SERVICE_CONFIG)

    def generate(self, variables: Dict[str, any]) -> Document:
        """Generate Service Agreement.

        Required variables:
        - service_provider: Service provider name
        - customer: Customer name
        - service_description: Description of services
        - start_date: Start date
        - jurisdiction: Governing law jurisdiction
        """
        self.variables = variables

        parties = [
            {"name": variables["service_provider"], "type": "Service Provider"},
            {"name": variables["customer"], "type": "Customer"},
        ]
        self._add_title_page("SERVICE AGREEMENT", parties, variables["start_date"])

        recitals = [
            f'{variables["customer"]} requires certain services to be provided',
            f'{variables["service_provider"]} is in the business of providing such services',
            "the parties desire to enter into an agreement for the provision of services",
        ]
        self._add_recitals(recitals)

        self._add_numbered_section(1, "Services")
        self._add_clause_with_substitution("Service Description", level=3, contract_type="service")
        self._add_clause_with_substitution("Service Standards", level=3, contract_type="service")

        self._add_numbered_section(1, "Fees and Payment")
        self._add_clause_with_substitution("Service Fees", level=3, contract_type="service")
        self._add_clause_with_substitution("Payment Schedule", level=3, contract_type="service")

        self._add_numbered_section(1, "Term and Renewal")
        self._add_clause_with_substitution("Term and Renewal", level=3, contract_type="service")

        self._add_numbered_section(1, "Warranties")
        self._add_clause_with_substitution("Warranties", level=3, contract_type="service")

        self._add_numbered_section(1, "General Provisions")
        self._add_numbered_section(2, "Governing Law")
        self._add_clause_with_substitution("Governing Law", level=3)
        self._add_numbered_section(2, "Entire Agreement")
        self._add_clause_with_substitution("Entire Agreement", level=3)

        self._add_signature_block(parties)

        return self.doc


# Template registry
TEMPLATES = {
    "saas": SaaSAgreementTemplate,
    "employment": EmploymentAgreementTemplate,
    "nda": NDATemplate,
    "consulting": ConsultingAgreementTemplate,
    "service": ServiceAgreementTemplate,
}


def get_template(contract_type: str) -> BaseContractTemplate:
    """Get template instance for a contract type.

    Args:
        contract_type: Contract type key (saas, employment, nda, consulting, service)

    Returns:
        Template instance

    Raises:
        ValueError: If contract type is invalid
    """
    template_class = TEMPLATES.get(contract_type)
    if not template_class:
        raise ValueError(f"Unknown contract type: {contract_type}")
    return template_class()
