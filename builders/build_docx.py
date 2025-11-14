"""DOCX→PDF pipeline for synthetic corpus generation.

This module provides functionality to:
1. Render DOCX templates with test data (python-docx-template)
2. Convert rendered DOCX to PDF (LibreOffice headless)
3. Extract ground truth from DOCX structure
4. Calculate and verify PDF hashes
5. Update metadata YAML files

Dependencies:
- docxtpl (python-docx-template)
- LibreOffice (soffice command)
- python-docx (for GT extraction)
"""

import hashlib
import subprocess
from pathlib import Path
from typing import Any

import yaml
from docx import Document
from docxtpl import DocxTemplate  # type: ignore[import-untyped]


class DocxBuilder:
    """Build PDF documents from DOCX templates with ground truth extraction."""

    def __init__(self, templates_dir: Path, output_dir: Path) -> None:
        """Initialize DocxBuilder.

        Args:
            templates_dir: Directory containing DOCX templates
            output_dir: Directory for output files (.docx, .pdf, .yaml)
        """
        self.templates_dir = templates_dir
        self.output_dir = output_dir
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def render_template(
        self, template_name: str, context: dict[str, Any], output_name: str
    ) -> Path:
        """Render DOCX template with context data.

        Args:
            template_name: Name of template file in templates_dir
            context: Dictionary of variables for Jinja2 substitution
            output_name: Name for output DOCX file (without extension)

        Returns:
            Path to rendered DOCX file

        Raises:
            FileNotFoundError: Template file not found
            RuntimeError: Template rendering failed
        """
        template_path = self.templates_dir / template_name
        if not template_path.exists():
            raise FileNotFoundError(f"Template not found: {template_path}")

        output_path = self.output_dir / f"{output_name}.docx"

        try:
            doc = DocxTemplate(str(template_path))
            doc.render(context)
            doc.save(str(output_path))
        except Exception as e:
            raise RuntimeError(f"Failed to render template: {e}") from e

        return output_path

    def convert_to_pdf(self, docx_path: Path) -> Path:
        """Convert DOCX to PDF using LibreOffice headless.

        Args:
            docx_path: Path to DOCX file to convert

        Returns:
            Path to generated PDF file

        Raises:
            FileNotFoundError: DOCX file not found
            RuntimeError: Conversion failed
        """
        if not docx_path.exists():
            raise FileNotFoundError(f"DOCX file not found: {docx_path}")

        # LibreOffice command: convert to PDF in the same directory
        cmd = [
            "soffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(docx_path.parent),
            str(docx_path),
        ]

        try:
            subprocess.run(
                cmd,
                capture_output=True,
                text=True,
                timeout=60,  # 60 second timeout
                check=True,
            )
        except subprocess.TimeoutExpired as e:
            raise RuntimeError(
                "LibreOffice conversion timeout after 60s"
            ) from e
        except subprocess.CalledProcessError as e:
            raise RuntimeError(
                f"LibreOffice conversion failed: {e.stderr}"
            ) from e

        # Check PDF was created
        pdf_path = docx_path.with_suffix(".pdf")
        if not pdf_path.exists():
            raise RuntimeError(
                f"PDF not created by LibreOffice: {pdf_path}"
            )

        return pdf_path

    def extract_ground_truth(self, docx_path: Path) -> str:
        """Extract ground truth Markdown from rendered DOCX structure.

        Args:
            docx_path: Path to rendered DOCX file

        Returns:
            Ground truth Markdown string

        Raises:
            FileNotFoundError: DOCX file not found
        """
        if not docx_path.exists():
            raise FileNotFoundError(f"DOCX file not found: {docx_path}")

        doc = Document(str(docx_path))
        markdown_lines: list[str] = []

        for element in doc.element.body:
            # Handle paragraphs
            if element.tag.endswith("p"):
                para = None
                for p in doc.paragraphs:
                    if p._element == element:
                        para = p
                        break

                if para and para.style:
                    # Check if it's a heading
                    if para.style.name and para.style.name.startswith("Heading"):
                        level = int(para.style.name.split()[-1])
                        markdown_lines.append(f"{'#' * level} {para.text}")
                    else:
                        # Regular paragraph
                        text = para.text.strip()
                        if text:
                            markdown_lines.append(text)

            # Handle tables
            elif element.tag.endswith("tbl"):
                table = None
                for t in doc.tables:
                    if t._element == element:
                        table = t
                        break

                if table:
                    markdown_lines.append(
                        self._table_to_markdown(table)
                    )

        return "\n\n".join(markdown_lines)

    def _table_to_markdown(self, table: Any) -> str:
        """Convert DOCX table to Markdown (GFM format).

        Args:
            table: python-docx table object

        Returns:
            Markdown table string
        """
        if not table.rows:
            return ""

        lines: list[str] = []

        # Header row
        header_cells = [cell.text.strip() for cell in table.rows[0].cells]
        lines.append("| " + " | ".join(header_cells) + " |")

        # Separator row
        lines.append("| " + " | ".join(["---"] * len(header_cells)) + " |")

        # Data rows
        for row in table.rows[1:]:
            cells = [cell.text.strip() for cell in row.cells]
            lines.append("| " + " | ".join(cells) + " |")

        return "\n".join(lines)

    def calculate_hash(self, file_path: Path) -> str:
        """Calculate SHA256 hash of file.

        Args:
            file_path: Path to file

        Returns:
            Hex digest of SHA256 hash

        Raises:
            FileNotFoundError: File not found
        """
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        sha256 = hashlib.sha256()
        with open(file_path, "rb") as f:
            for chunk in iter(lambda: f.read(8192), b""):
                sha256.update(chunk)

        return sha256.hexdigest()

    def create_metadata(
        self,
        output_name: str,
        title: str,
        source: str,
        pdf_hash: str,
        docx_hash: str,
        context: dict[str, Any],
    ) -> Path:
        """Create metadata YAML file for generated document.

        Args:
            output_name: Base name for output files
            title: Document title
            source: Source template name
            pdf_hash: SHA256 hash of PDF
            docx_hash: SHA256 hash of DOCX
            context: Template context used for rendering

        Returns:
            Path to metadata YAML file
        """
        metadata = {
            "title": title,
            "source": source,
            "corpus": "synthetic",
            "category": "docx_generated",
            "license": "CC0-1.0",
            "generation": {
                "method": "docx_template",
                "template": source,
                "builder": "python-docx-template + LibreOffice",
                "context_variables": len(context),
            },
            "files": {
                "docx": f"{output_name}.docx",
                "pdf": f"{output_name}.pdf",
                "ground_truth": f"{output_name}.md",
            },
            "hashes": {
                "pdf_sha256": pdf_hash,
                "docx_sha256": docx_hash,
            },
        }

        yaml_path = self.output_dir / f"{output_name}.yaml"
        with open(yaml_path, "w") as f:
            yaml.dump(metadata, f, sort_keys=False, default_flow_style=False)

        return yaml_path

    def build_document(
        self,
        template_name: str,
        output_name: str,
        title: str,
        context: dict[str, Any],
    ) -> dict[str, Path]:
        """Full pipeline: render template, convert to PDF, extract GT, create metadata.

        Args:
            template_name: Name of template file
            output_name: Base name for output files
            title: Document title for metadata
            context: Template context variables

        Returns:
            Dictionary with paths: docx, pdf, ground_truth, metadata

        Raises:
            Various exceptions from pipeline steps
        """
        print(f"Building {output_name} from {template_name}...")

        # Step 1: Render template
        print("  1/5 Rendering template...")
        docx_path = self.render_template(template_name, context, output_name)
        print(f"      → {docx_path.name}")

        # Step 2: Convert to PDF
        print("  2/5 Converting to PDF (LibreOffice)...")
        pdf_path = self.convert_to_pdf(docx_path)
        print(f"      → {pdf_path.name}")

        # Step 3: Extract ground truth
        print("  3/5 Extracting ground truth...")
        ground_truth = self.extract_ground_truth(docx_path)
        gt_path = self.output_dir / f"{output_name}.md"
        with open(gt_path, "w") as f:
            f.write(ground_truth)
        print(f"      → {gt_path.name}")

        # Step 4: Calculate hashes
        print("  4/5 Calculating hashes...")
        pdf_hash = self.calculate_hash(pdf_path)
        docx_hash = self.calculate_hash(docx_path)
        print(f"      PDF:  {pdf_hash[:16]}...")
        print(f"      DOCX: {docx_hash[:16]}...")

        # Step 5: Create metadata
        print("  5/5 Creating metadata...")
        yaml_path = self.create_metadata(
            output_name, title, template_name, pdf_hash, docx_hash, context
        )
        print(f"      → {yaml_path.name}")

        print(f"✓ Complete: {output_name}")
        print()

        return {
            "docx": docx_path,
            "pdf": pdf_path,
            "ground_truth": gt_path,
            "metadata": yaml_path,
        }


def main() -> None:
    """Example usage and testing."""
    templates_dir = Path("corpus/synthetic/templates")
    output_dir = Path("corpus/synthetic")

    builder = DocxBuilder(templates_dir, output_dir)

    # Example: simple report
    simple_context = {
        "title": "Test Simple Report",
        "author": "PDF-Bench Test Suite",
        "date": "2025-11-14",
        "report_id": "TST-001",
        "executive_summary": "This is a test document.",
        "introduction": "Test introduction paragraph.",
        "background": "Test background information.",
        "methodology": "Test methodology description.",
        "findings": "Test findings section.",
        "observations": ["First observation", "Second observation"],
        "results": "Test results section.",
        "conclusion": "Test conclusion paragraph.",
        "recommendations": ["Recommendation 1", "Recommendation 2"],
        "references": "Test references section.",
    }

    try:
        result = builder.build_document(
            template_name="01_simple_report.docx",
            output_name="test_simple_report",
            title="Test Simple Report",
            context=simple_context,
        )

        print("Generated files:")
        for key, path in result.items():
            print(f"  {key}: {path}")

    except Exception as e:
        print(f"Error: {e}")
        raise


if __name__ == "__main__":
    main()
